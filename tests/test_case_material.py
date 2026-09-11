from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pytest
from docx import Document

from agents.collector import EvidenceCollector
from case_file import CaseFile
from case_material import (
    CaseFileCurator,
    CaseMaterialBlockKind,
    CaseMaterialCuratorError,
    CaseMaterialInput,
)
from orchestrator import InvestigationEventKind, stream_investigation


def _docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Observatory report", level=1)
    document.add_paragraph("The lens cabinet was found open.")
    document.add_paragraph("Check the roof hatch.", style="List Number")
    document.add_paragraph("A brass key was missing.", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Time"
    table.cell(0, 1).text = "Event"
    table.cell(1, 0).text = "21:00"
    table.cell(1, 1).text = "Cabinet opened"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def test_curator_preserves_docx_structure_and_table_cell_locations() -> None:
    result = CaseFileCurator().curate(
        [CaseMaterialInput(display_name="/temporary/uploads/report.docx", content=_docx_bytes())]
    )

    assert [block.kind for block in result.blocks] == [
        CaseMaterialBlockKind.SECTION,
        CaseMaterialBlockKind.PARAGRAPH,
        CaseMaterialBlockKind.ORDERED_LIST_ITEM,
        CaseMaterialBlockKind.UNORDERED_LIST_ITEM,
        CaseMaterialBlockKind.TABLE,
    ]
    table = result.blocks[-1].table
    assert table is not None
    assert table.headers == ("Time", "Event")
    assert table.rows == (("21:00", "Cabinet opened"),)
    assert table.cell_references[0][1].table_row == 1
    assert table.cell_references[0][1].table_column == 2
    assert table.header_references[1].table_column == 2
    assert result.blocks[-1].source_reference.source_name == "report.docx"


def test_curator_excludes_bad_docx_but_keeps_usable_companion_material() -> None:
    result = CaseFileCurator().curate(
        [
            CaseMaterialInput(display_name="broken.docx", content=b"not a document"),
            CaseMaterialInput(display_name="case.txt", content="A telescope vanished."),
        ]
    )

    assert result.canonical_text == "A telescope vanished."
    assert any("broken.docx" in warning and "could not be read" in warning.lower() for warning in result.warnings)


def test_docx_follows_pasted_material_in_the_supplied_source_order() -> None:
    result = CaseFileCurator().curate(
        [
            CaseMaterialInput(display_name="Pasted case text", content="The dome alarm sounded."),
            CaseMaterialInput(display_name="report.docx", content=_docx_bytes()),
        ]
    )

    assert list(result.source_text) == ["Pasted case text", "report.docx"]
    assert result.blocks[0].id == "Pasted case text:S-001"
    assert result.blocks[1].id == "report.docx:S-001"


def test_corrupt_docx_halts_before_investigation_when_no_material_is_usable() -> None:
    events = list(
        stream_investigation(
            "", TableRowCitingLLM(), [CaseMaterialInput(display_name="broken.docx", content=b"bad")]
        )
    )

    assert [event.kind for event in events] == [
        InvestigationEventKind.CASE_MATERIAL_CURATION_STARTED,
        InvestigationEventKind.VALIDATION_ERROR,
    ]
    assert "No usable case material" in (events[-1].message or "")
    assert "could not be read" in (events[-1].message or "").lower()


def test_curator_reports_encrypted_and_empty_docx_files_without_inventing_material() -> None:
    empty_document = Document()
    empty_content = BytesIO()
    empty_document.save(empty_content)

    result = CaseFileCurator().curate(
        [
            CaseMaterialInput(display_name="locked.docx", content=b"\xd0\xcf\x11\xe0encrypted"),
            CaseMaterialInput(display_name="empty.docx", content=empty_content.getvalue()),
            CaseMaterialInput(display_name="case.txt", content="The map disappeared."),
        ]
    )

    assert result.canonical_text == "The map disappeared."
    assert any("locked.docx" in warning and "encrypted" in warning.lower() for warning in result.warnings)
    assert any("empty.docx" in warning and "no usable case text" in warning.lower() for warning in result.warnings)


def test_curator_preserves_markdown_structure_and_safe_source_locations() -> None:
    result = CaseFileCurator().curate(
        [
            CaseMaterialInput(
                display_name="/temporary/uploads/incident.md",
                content=(
                    "# Gallery report\n\nThe alarm sounded.\n\n"
                    "1. Lock the east door\n2. Call the curator\n\n"
                    "- A muddy print\n\n"
                    "| Time | Event |\n| --- | --- |\n| 21:00 | Alarm |\n"
                ),
            )
        ]
    )

    assert [block.kind for block in result.blocks] == [
        CaseMaterialBlockKind.SECTION,
        CaseMaterialBlockKind.PARAGRAPH,
        CaseMaterialBlockKind.ORDERED_LIST_ITEM,
        CaseMaterialBlockKind.ORDERED_LIST_ITEM,
        CaseMaterialBlockKind.UNORDERED_LIST_ITEM,
        CaseMaterialBlockKind.TABLE,
    ]
    assert result.blocks[2].ordinal == 1
    assert result.blocks[-1].table is not None
    assert result.blocks[-1].table.headers == ("Time", "Event")
    assert result.blocks[-1].table.rows == (("21:00", "Alarm"),)
    assert result.blocks[-1].source_reference.source_name == "incident.md"
    assert "/" not in result.blocks[-1].source_reference.source_name
    assert "# Gallery report" in result.canonical_text


def test_curator_keeps_pasted_and_markdown_sources_in_supplied_order() -> None:
    result = CaseFileCurator().curate(
        [
            CaseMaterialInput(display_name="Pasted case text", content="First report."),
            CaseMaterialInput(display_name="notes.md", content="# Second report"),
            CaseMaterialInput(display_name="notes.md", content="# Third report"),
        ]
    )

    assert list(result.source_text) == ["Pasted case text", "notes.md", "notes (3).md"]
    assert [block.id for block in result.blocks] == [
        "Pasted case text:S-001",
        "notes.md:S-001",
        "notes (3).md:S-001",
    ]


def test_curator_halts_with_a_safe_warning_when_every_source_is_unsupported() -> None:
    with pytest.raises(CaseMaterialCuratorError, match="No usable case material") as error:
        CaseFileCurator().curate(
            [CaseMaterialInput(display_name="evidence.csv", content=b"not a CSV")]
        )

    assert "evidence.csv" in str(error.value)
    assert "Unsupported file type" in str(error.value)
    assert "/" not in str(error.value)


def test_curator_excludes_sealed_facilitator_content_from_canonical_material() -> None:
    result = CaseFileCurator().curate(
        [
            CaseMaterialInput(
                display_name="workshop.txt",
                content=(
                    "Participant briefing: the lantern disappeared.\n\n"
                    "Facilitator Only Solution\nThe answer is withheld."
                ),
            )
        ]
    )

    assert "Participant briefing" in result.canonical_text
    assert "answer is withheld" not in result.canonical_text
    assert "answer is withheld" not in result.source_text["workshop.txt"]
    assert any("sealed facilitator" in warning.lower() for warning in result.warnings)


def test_docx_keeps_participant_prose_that_mentions_a_sealed_solution_before_a_heading() -> None:
    document = Document()
    document.add_heading("Participant materials", level=1)
    document.add_paragraph("This edition is followed by a sealed facilitator solution.")
    document.add_paragraph("The observatory key was missing.")
    document.add_heading("Facilitator Only Solution", level=1)
    document.add_paragraph("Sealed evaluation conclusion.")
    content = BytesIO()
    document.save(content)

    result = CaseFileCurator().curate(
        [CaseMaterialInput(display_name="workshop.docx", content=content.getvalue())]
    )

    assert "sealed facilitator solution" in result.canonical_text
    assert "observatory key was missing" in result.canonical_text
    assert "Facilitator Only Solution" not in result.canonical_text
    assert "Sealed evaluation conclusion" not in result.canonical_text
    assert [block.id for block in result.blocks] == [
        "workshop.docx:S-001",
        "workshop.docx:S-002",
        "workshop.docx:S-003",
    ]
    assert len([warning for warning in result.warnings if "sealed facilitator" in warning]) == 1


def test_markdown_keeps_participant_prose_that_mentions_a_sealed_solution_before_a_heading() -> None:
    result = CaseFileCurator().curate(
        [
            CaseMaterialInput(
                display_name="workshop.md",
                content=(
                    "# Participant materials\n\n"
                    "This edition is followed by a sealed facilitator solution.\n\n"
                    "The observatory key was missing.\n\n"
                    "# Facilitator Evaluation Material\n\n"
                    "Sealed evaluation conclusion."
                ),
            )
        ]
    )

    assert "sealed facilitator solution" in result.canonical_text
    assert "observatory key was missing" in result.canonical_text
    assert "Facilitator Evaluation Material" not in result.canonical_text
    assert "Sealed evaluation conclusion" not in result.canonical_text
    assert [block.id for block in result.blocks] == [
        "workshop.md:S-001",
        "workshop.md:S-002",
        "workshop.md:S-003",
    ]
    assert len([warning for warning in result.warnings if "sealed facilitator" in warning]) == 1


def test_reference_case_docx_retains_participant_material_and_excludes_the_sealed_section() -> None:
    reference_case = Path("docs/reference-material/The_Vanishing_Aurora_Diamond_Case_Book.docx")

    result = CaseFileCurator().curate(
        [CaseMaterialInput(display_name=reference_case.name, content=reference_case.read_bytes())]
    )

    assert "Participant edition followed by a sealed facilitator solution" in result.canonical_text
    assert "Evidence Register" in result.canonical_text
    assert any(block.table is not None for block in result.blocks)
    assert "Facilitator Only Solution" not in result.canonical_text
    assert "Most likely explanation" not in result.canonical_text
    assert len([warning for warning in result.warnings if "sealed facilitator" in warning]) == 1


class ReferenceParticipantLLM:
    def __init__(self) -> None:
        self.prompts: list[str] = []

    def call_llm(self, prompt: str, system: str, response_schema: dict) -> dict:
        self.prompts.append(prompt)
        return {
            "evidence": [
                {
                    "id": "E-01",
                    "statement": "An access record appears in the evidence register.",
                    "classification": "observed_fact",
                    "source_reference_ids": [
                        "The_Vanishing_Aurora_Diamond_Case_Book.docx:S-060:R-001"
                    ],
                }
            ]
        }


def test_reference_case_docx_excludes_sealed_material_before_evidence_collection() -> None:
    reference_case = Path("docs/reference-material/The_Vanishing_Aurora_Diamond_Case_Book.docx")
    llm = ReferenceParticipantLLM()

    events = stream_investigation(
        "",
        llm,
        [CaseMaterialInput(display_name=reference_case.name, content=reference_case.read_bytes())],
    )
    first_events = [next(events) for _ in range(4)]

    assert [event.kind for event in first_events] == [
        InvestigationEventKind.CASE_MATERIAL_CURATION_STARTED,
        InvestigationEventKind.CASE_MATERIAL_CURATION_COMPLETED,
        InvestigationEventKind.EVIDENCE_COLLECTION_STARTED,
        InvestigationEventKind.EVIDENCE_COLLECTION_COMPLETED,
    ]
    case_file = first_events[-1].case_file
    assert case_file is not None
    assert "Participant edition followed by a sealed facilitator solution" in case_file.mystery_text
    assert "At 8:00 PM, curator Dr. Mira Sen displayed the Aurora Diamond" in case_file.mystery_text
    assert "Evidence Register" in case_file.mystery_text
    assert any(block.table is not None for block in case_file.material_blocks)
    assert "Facilitator Only Solution" not in case_file.canonical_material
    assert "Most likely explanation" not in case_file.canonical_material
    assert all("Facilitator Only Solution" not in source for source in case_file.source_text.values())
    assert all("Most likely explanation" not in source for source in case_file.source_text.values())
    assert all("Facilitator Only Solution" not in prompt for prompt in llm.prompts)
    assert all("Most likely explanation" not in prompt for prompt in llm.prompts)
    assert case_file.evidence[0].source_references[0].table_row == 1


def test_sealed_only_docx_halts_before_evidence_collection_without_exposing_its_content() -> None:
    document = Document()
    document.add_heading("Facilitator Only Solution", level=1)
    document.add_paragraph("Sealed evaluation conclusion.")
    content = BytesIO()
    document.save(content)

    with pytest.raises(CaseMaterialCuratorError, match="No usable case material") as error:
        CaseFileCurator().curate(
            [CaseMaterialInput(display_name="sealed.docx", content=content.getvalue())]
        )

    assert "excluded sealed facilitator solution" in str(error.value)
    assert "Sealed evaluation conclusion" not in str(error.value)


@pytest.mark.parametrize(
    "marker",
    ["Facilitator-only solution", "Facilitator evaluation material"],
)
def test_curator_recognizes_facilitator_section_variants(marker: str) -> None:
    result = CaseFileCurator().curate(
        [CaseMaterialInput(display_name="workshop.txt", content=f"Participant text.\n\n{marker}\nHidden.")]
    )

    assert "Hidden." not in result.canonical_text


class TableRowCitingLLM:
    def call_llm(self, prompt: str, system: str, response_schema: dict) -> dict:
        assert "[access.md:S-001:R-001]" in prompt
        return {
            "evidence": [
                {
                    "id": "E-01",
                    "statement": "The archive opened at 21:00.",
                    "classification": "observed_fact",
                    "source_reference_ids": ["access.md:S-001:R-001"],
                }
            ]
        }


def test_evidence_can_retain_a_specific_table_row_source_reference() -> None:
    curated = CaseFileCurator().curate(
        [
            CaseMaterialInput(
                display_name="access.md",
                content="| Time | Door |\n| --- | --- |\n| 21:00 | Archive |",
            )
        ]
    )
    case_file = CaseFile(
        mystery_text=curated.canonical_text,
        canonical_material=curated.prompt_text,
        material_blocks=curated.blocks,
    )

    EvidenceCollector(TableRowCitingLLM()).run(case_file)

    assert case_file.evidence[0].source_references[0].table_row == 1


class TableCellCitingLLM:
    def call_llm(self, prompt: str, system: str, response_schema: dict) -> dict:
        assert "[report.docx:S-005:R-001]" in prompt
        return {
            "evidence": [
                {
                    "id": "E-01",
                    "statement": "The cabinet opened at 21:00.",
                    "classification": "observed_fact",
                    "source_reference_ids": ["report.docx:S-005:R-001:C-002"],
                }
            ]
        }


def test_docx_evidence_can_retain_a_specific_table_cell_source_reference() -> None:
    curated = CaseFileCurator().curate(
        [CaseMaterialInput(display_name="report.docx", content=_docx_bytes())]
    )
    case_file = CaseFile(
        mystery_text=curated.canonical_text,
        canonical_material=curated.prompt_text,
        material_blocks=curated.blocks,
    )

    EvidenceCollector(TableCellCitingLLM()).run(case_file)

    reference = case_file.evidence[0].source_references[0]
    assert (reference.table_row, reference.table_column) == (1, 2)


class SealedDocxLLM:
    def call_llm(self, prompt: str, system: str, response_schema: dict) -> dict:
        assert "Participant fact" in prompt
        assert "mentions a sealed facilitator solution" in prompt
        assert "sealed conclusion" not in prompt
        return {
            "evidence": [
                {
                    "id": "E-01",
                    "statement": "The vault door was open.",
                    "classification": "observed_fact",
                    "source_reference_ids": ["workshop.docx:S-001"],
                }
            ]
        }


def test_public_docx_flow_excludes_sealed_facilitator_material_from_agent_input() -> None:
    document = Document()
    document.add_paragraph("Participant fact: the vault door was open.")
    document.add_paragraph("This guide mentions a sealed facilitator solution after the activity.")
    document.add_heading("Facilitator Only Solution", level=1)
    document.add_paragraph("sealed conclusion")
    content = BytesIO()
    document.save(content)

    events = stream_investigation(
        "",
        SealedDocxLLM(),
        [CaseMaterialInput(display_name="workshop.docx", content=content.getvalue())],
    )
    first_events = [next(events) for _ in range(4)]

    case_file = first_events[-1].case_file
    assert case_file is not None
    assert "sealed conclusion" not in case_file.canonical_material
    assert any("sealed facilitator" in warning.lower() for warning in case_file.material_warnings)


class SourceCitingLLM:
    def call_llm(self, prompt: str, system: str, response_schema: dict) -> dict:
        if "Evidence Collector" in system:
            assert "[report.md:S-001]" in prompt
            return {
                "evidence": [
                    {
                        "id": "E-01",
                        "statement": "The lantern went out.",
                        "classification": "observed_fact",
                        "source_block_ids": ["report.md:S-001"],
                    }
                ]
            }
        raise RuntimeError("Stop after the Evidence Collector for this intake test.")


def test_uploaded_markdown_is_curated_before_evidence_collection() -> None:
    events = stream_investigation(
        "",
        SourceCitingLLM(),
        [CaseMaterialInput(display_name="report.md", content="# The lantern went out.")],
    )
    first_events = [next(events) for _ in range(4)]

    assert [event.kind for event in first_events[:3]] == [
        InvestigationEventKind.CASE_MATERIAL_CURATION_STARTED,
        InvestigationEventKind.CASE_MATERIAL_CURATION_COMPLETED,
        InvestigationEventKind.EVIDENCE_COLLECTION_STARTED,
    ]
    curated_case_file = first_events[-1].case_file
    assert curated_case_file is not None
    assert curated_case_file.source_text == {"report.md": "# The lantern went out."}
    assert curated_case_file.evidence[0].source_references[0].source_name == "report.md"


class UnknownSourceReferenceLLM:
    def call_llm(self, prompt: str, system: str, response_schema: dict) -> dict:
        return {
            "evidence": [
                {
                    "id": "E-01",
                    "statement": "The lantern went out.",
                    "classification": "observed_fact",
                    "source_reference_ids": ["report.md:S-004"],
                }
            ]
        }


def test_unknown_evidence_source_reference_is_a_visible_step_failure() -> None:
    events = list(
        stream_investigation(
            "",
            UnknownSourceReferenceLLM(),
            [CaseMaterialInput(display_name="report.md", content="# The lantern went out.")],
        )
    )

    assert [event.kind for event in events] == [
        InvestigationEventKind.CASE_MATERIAL_CURATION_STARTED,
        InvestigationEventKind.CASE_MATERIAL_CURATION_COMPLETED,
        InvestigationEventKind.EVIDENCE_COLLECTION_STARTED,
        InvestigationEventKind.STEP_FAILED,
    ]
    assert events[-1].agent_name == "Evidence Collector"
    assert "unknown source reference IDs" in (events[-1].message or "")
